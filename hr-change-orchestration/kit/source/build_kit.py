"""Builds the HR Change + Orchestration kit from blueprint_data.py.

    python kit/source/build_kit.py

Writes to kit/:
  HR Change Reference.xlsx        reference tables the flows read (edit here, not in the flows)
  HR Change Build Blueprint.xlsx  roadmap, sprint/gate tracker, every build step, test pack
  HR Change Blueprint.docx        the step-by-step blueprint to follow
  HR-Change-Requests.csv          creates the one SharePoint list
"""
import csv
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import blueprint_data as D  # noqa: E402

OUT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

PLUM, FOREST, CHAR, GREY, MINT, MINT_LIGHT, DAHLIA = "512241", "0C4C4C", "282F2D", "999999", "BEE3DC", "E8F6F3", "B73563"
XL_FONT = "Calibri"
thin = Side(style="thin", color=GREY)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=PLUM)
ALT_FILL = PatternFill("solid", fgColor=MINT_LIGHT)
INPUT_FILL = PatternFill("solid", fgColor=MINT)


# ====================================================================== Excel helpers
def new_wb(title):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = title
    wb.properties.creator = "People Ops"
    return wb


def brand_sheet(ws, doc_title):
    ws.oddHeader.left.text = "Greater Good Health"
    ws.oddHeader.center.text = doc_title
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_view.showGridLines = False


def title_block(ws, title, subtitle, row=1):
    ws.cell(row=row, column=1, value=title).font = Font(name=XL_FONT, size=16, bold=True, color=PLUM)
    ws.cell(row=row + 1, column=1, value=subtitle).font = Font(name=XL_FONT, size=11, italic=True, color=CHAR)
    return row + 3


def write_table(ws, top, headers, rows, name, widths=None, wrap=True, validations=None, formulas=False):
    """Write a formatted Excel table starting at row `top`. Returns the last row."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font = Font(name=XL_FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=top + r, column=c, value=v)
            if isinstance(v, str) and v.startswith("=") and not formulas:
                cell.data_type = "s"  # SharePoint formula text, not an Excel formula
                cell.quotePrefix = True
            cell.font = Font(name=XL_FONT, color=CHAR)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if r % 2 == 1:
                cell.fill = ALT_FILL
    last = top + max(len(rows), 1)
    ref = f"A{top}:{get_column_letter(len(headers))}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    for col_name, options in (validations or {}).items():
        idx = headers.index(col_name) + 1
        col = get_column_letter(idx)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Pick a value from the list.", "Not allowed"
        ws.add_data_validation(dv)
        dv.add(f"{col}{top + 1}:{col}{max(last, top + 300)}")
    return last


def note_rows(ws, start, lines, col=1):
    for i, line in enumerate(lines):
        c = ws.cell(row=start + i, column=col, value=line)
        c.font = Font(name=XL_FONT, color=CHAR, bold=line.endswith(":"))
        c.alignment = Alignment(wrap_text=False)
    return start + len(lines)


# ====================================================================== 1. Reference workbook
def build_reference():
    title = "HR Change Reference"
    wb = new_wb(title)

    ws = wb.create_sheet("About")
    brand_sheet(ws, title)
    r = title_block(ws, "HR Change Reference — tables the flows read", "Edit SLAs, approval chains, checklists, approvers and settings here. No flow edits needed.")
    r = note_rows(ws, r, [
        "Where this file lives:",
        f"  People Ops SharePoint site → Documents → HR Change → {D.REF_FILE}  (HR owners edit; nobody else needs access).",
        "",
        "The contract (break these and the flows stop):",
        "  1. Never rename a table: tblSettings, tblChangeTypes, tblChecklist, tblApprovers.",
        "  2. Never rename, delete or reorder-rename a column header in those tables. Adding rows is always fine.",
        "  3. Add rows INSIDE the table (Tab from the last cell, or type in the row right under it).",
        "  4. Yes/No columns contain exactly Yes or No. Step is two digits (01, 02 … 16) so steps sort correctly.",
        "  5. Code must be the change type's prefix (A1, A2, B … I, Other) — the flows read it from the start of the Change Type value.",
        "  6. ApprovalChain lists roles in order, separated by '>':  Manager > Regional > HR.  Roles: Manager, Regional, HR, Payroll. Blank = no approval.",
        "",
        "What each table drives:",
        "  tblSettings     → prefix, time zone, payroll cutoffs, reminder timing, front-door link, admin email (all flows)",
        "  tblChangeTypes  → SLA, due-date rule, approval chain, default assignee, close-the-loop email, People Ops summary, sensitive (Flows 1, 2, 4, 5, 7)",
        "  tblChecklist    → the checklist posted to each new request (Flow 2); Auto = 'approval' / 'resolution' steps tick themselves (Flows 4, 5)",
        "  tblApprovers    → Regional / HR / Payroll approvers by region; Manager comes from the request (Flow 5)",
        "",
        "Reference only (not read by flows): List Fields, Statuses, Views — use them to build and maintain the SharePoint list.",
        "Changes take effect for the NEXT request submitted. Existing requests keep the checklist they were given.",
    ])
    ws.column_dimensions["A"].width = 150

    # Settings: single-row table + guide
    ws = wb.create_sheet("Settings")
    brand_sheet(ws, title)
    r = title_block(ws, "Settings", "One row. Edit the values; keep the headers.")
    keys = list(D.SETTINGS)
    last = write_table(ws, r, keys, [[D.SETTINGS[k][0] for k in keys]], "tblSettings", widths=[16, 24, 12, 12, 22, 60, 26])
    for c in range(1, len(keys) + 1):
        ws.cell(row=last, column=c).fill = INPUT_FILL
    g = last + 3
    ws.cell(row=g - 1, column=1, value="Setting guide").font = Font(name=XL_FONT, bold=True, color=FOREST, size=12)
    write_table(ws, g, ["Setting", "What it does"], [[k, D.SETTINGS[k][1]] for k in keys], "refSettingGuide")
    ws.freeze_panes = ws.cell(row=r + 1, column=1)

    ws = wb.create_sheet("Change Types")
    brand_sheet(ws, title)
    r = title_block(ws, "Change Types", "One row per change type. SLA, approval chain and notifications.")
    write_table(ws, r, D.CHANGE_TYPE_COLS, [list(c) for c in D.CHANGE_TYPES], "tblChangeTypes",
                widths=[8, 36, 12, 18, 30, 26, 30, 14, 18, 11, 48, 60],
                validations={"DueDateRule": ["Business days", "Effective date", "None"], "CloseLoopEmail": ["Yes", "No"],
                             "PostSummaryToPeopleOps": ["Yes", "No"], "Sensitive": ["Yes", "No"]})
    ws.freeze_panes = ws.cell(row=r + 1, column=3)

    ws = wb.create_sheet("Checklists")
    brand_sheet(ws, title)
    r = title_block(ws, "Checklists (master SOP templates)", "Standing reference — never completed. Flow 2 copies the active steps into each new request.")
    write_table(ws, r, D.CHECKLIST_COLS, [[s[c] for c in D.CHECKLIST_COLS] for s in D.CHECKLISTS], "tblChecklist",
                widths=[8, 7, 60, 70, 10, 11, 14, 14, 8],
                validations={"Code": [c[0] for c in D.CHANGE_TYPES], "Auto": ["approval", "resolution", "jira"],
                             "AfterApproval": ["Yes", "No"], "DueOnLastDay": ["Yes", "No"], "Active": ["Yes", "No"],
                             "Owner": ["HR", "CJ", "Payroll", "IT", "Talent", "Manager"]})
    for row in ws.iter_rows(min_row=r + 1, max_row=r + len(D.CHECKLISTS), min_col=2, max_col=2):
        for c in row:
            c.number_format = "@"
    ws.freeze_panes = ws.cell(row=r + 1, column=3)

    ws = wb.create_sheet("Approvers")
    brand_sheet(ws, title)
    r = title_block(ws, "Approvers", "Regional / HR / Payroll approvers. Region '*' = any region. Manager approvals use the request's Manager field.")
    write_table(ws, r, D.APPROVER_COLS, [list(a) for a in D.APPROVERS], "tblApprovers", widths=[12, 12, 36, 28, 8, 80],
                validations={"Role": ["Regional", "HR", "Payroll"], "Active": ["Yes", "No"]})
    ws.freeze_panes = ws.cell(row=r + 1, column=1)

    ws = wb.create_sheet("List Fields")
    brand_sheet(ws, title)
    r = title_block(ws, f"List Fields — the one list: {D.LIST_NAME}", "Reference for building the list. Internal names (no spaces) are what the flows use.")
    rows = [[f[0], f[1], f[2], f[3], f[4], f[5], f[6], D.show_formula(f), f[7]] for f in D.FIELDS]
    write_table(ws, r, ["Internal name", "Display name", "Type", "Required", "Choices / default", "Group", "Show for codes", "Show formula (form)", "Set by"],
                rows, "refListFields", widths=[22, 28, 24, 10, 60, 12, 12, 90, 50])
    ws.freeze_panes = ws.cell(row=r + 1, column=2)

    ws = wb.create_sheet("Statuses")
    brand_sheet(ws, title)
    r = title_block(ws, "Statuses (Tag, You're It)", "Unchanged from the plan, plus Waiting On / Next Action.")
    write_table(ws, r, ["Status", "Meaning", "Set by", "What happens next"], [list(s) for s in D.STATUSES], "refStatuses", widths=[14, 36, 60, 60])

    ws = wb.create_sheet("Views")
    brand_sheet(ws, title)
    r = title_block(ws, "Workbench views", "Create these on the list (All Items → Create new view).")
    write_table(ws, r, ["View", "Filter", "Sort / group", "Columns", "Purpose"], [list(v) for v in D.VIEWS], "refViews", widths=[24, 46, 34, 60, 46])

    path = os.path.join(OUT, D.REF_FILE)
    wb.save(path)
    return path


# ====================================================================== 2. Blueprint tracker workbook
STATUS_OPTS = ["Not started", "In progress", "Done", "Blocked", "Skipped"]


def paste_text(paste):
    return "\n".join(f"{label}:\n    {text}" for label, text in paste)


def build_tracker():
    title = "HR Change Build Blueprint"
    wb = new_wb(title)

    # Overview with live progress formulas
    ws = wb.create_sheet("Overview")
    brand_sheet(ws, title)
    title_block(ws, "HR Change + Orchestration — Build Blueprint", "Sprints, gates and progress. Mint cells are inputs; everything else is calculated.")
    ws["A4"], ws["B4"] = "Sprint 0 start date", "=TODAY()"
    ws["A4"].font = Font(name=XL_FONT, bold=True, color=CHAR)
    ws["B4"].fill = INPUT_FILL
    ws["B4"].number_format = "mmm d, yyyy"
    ws["C4"] = "← type your real start date (a Monday)"
    ws["C4"].font = Font(name=XL_FONT, italic=True, color=CHAR)
    hdr = ["Sprint", "Name", "Phase", "Weeks", "Start", "End", "Steps", "Done", "% done", "Gate", "Gate status"]
    top = 6
    rows = []
    for i, s in enumerate(D.SPRINTS):
        r = top + 1 + i
        start = "=B4" if i == 0 else f"=IF(F{r - 1}=\"\",\"\",F{r - 1}+1)"
        end = f"=IF(OR(E{r}=\"\",D{r}=0),\"\",E{r}+D{r}*7-1)"
        rows.append([s[0], s[1], s[3], s[2], start, end,
                     f"=COUNTIF('Build Steps'!$B:$B,A{r})",
                     f"=COUNTIFS('Build Steps'!$B:$B,A{r},'Build Steps'!$J:$J,\"Done\")",
                     f"=IF(G{r}=0,\"\",H{r}/G{r})", s[6], "Not started"])
    last = write_table(ws, top, hdr, rows, "tblSprintSummary", widths=[9, 46, 9, 8, 13, 13, 8, 8, 9, 26, 14], formulas=True,
                       validations={"Gate status": ["Not started", "In progress", "Passed", "Blocked"]})
    for r in range(top + 1, last + 1):
        ws[f"E{r}"].number_format = ws[f"F{r}"].number_format = "mmm d, yyyy"
        ws[f"I{r}"].number_format = "0%"
        ws[f"K{r}"].fill = INPUT_FILL
    tr = last + 1
    ws[f"A{tr}"] = "Total"
    ws[f"G{tr}"] = f"=SUM(G{top + 1}:G{last})"
    ws[f"H{tr}"] = f"=SUM(H{top + 1}:H{last})"
    ws[f"I{tr}"] = f"=IF(G{tr}=0,\"\",H{tr}/G{tr})"
    ws[f"I{tr}"].number_format = "0%"
    for c in "AGHI":
        ws[f"{c}{tr}"].font = Font(name=XL_FONT, bold=True, color=CHAR)
    note_rows(ws, tr + 2, [
        "How to use this workbook:",
        "  1. Set the start date above. Sprint dates follow (S6 is ongoing).",
        "  2. Work the 'Build Steps' tab top to bottom. Set Status as you go — progress updates here.",
        "  3. Each sprint ends at a gate (Gates tab). Don't start the next sprint until every gate line is Passed.",
        "  4. Run the tests named in each sprint (Test Pack tab). Record Pass/Fail.",
        "  5. The Word document 'HR Change Blueprint' has the same steps with full explanations.",
    ])

    ws = wb.create_sheet("Roadmap")
    brand_sheet(ws, title)
    r = title_block(ws, "Roadmap — sprints and gates", "Crawl → Walk → Run, as in the project plan.")
    rows = []
    for i, s in enumerate(D.SPRINTS):
        o = 7 + i
        rows.append([s[0], s[1], s[3], f"=Overview!E{o}", f"=Overview!F{o}", s[4], s[5], s[6], "\n".join("• " + g for g in s[7])])
    last = write_table(ws, r, ["Sprint", "Name", "Phase", "Start", "End", "Goal", "Entry criteria", "Gate", "Gate criteria"], rows, "tblRoadmap", formulas=True,
                       widths=[8, 40, 9, 13, 13, 60, 30, 24, 70])
    for rr in range(r + 1, last + 1):
        ws[f"D{rr}"].number_format = ws[f"E{rr}"].number_format = "mmm d, yyyy"
    ws.freeze_panes = ws.cell(row=r + 1, column=3)

    ws = wb.create_sheet("Build Steps")
    brand_sheet(ws, title)
    r = title_block(ws, "Build Steps", "Every step in order. Status is the only column you need to update (plus Done on / Notes).")
    rows = [[s["id"], s["sprint"], s["area"], s["title"], s["where"], "\n".join("• " + x for x in s["do"]), paste_text(s["paste"]),
             s["check"], s["owner"], "Not started", None, ""] for s in D.STEPS]
    last = write_table(ws, r, ["ID", "Sprint", "Area", "Step", "Where", "Do", "Paste / settings", "Check", "Owner", "Status", "Done on", "Notes"],
                       rows, "tblBuildSteps", widths=[8, 7, 13, 38, 34, 60, 90, 40, 10, 13, 12, 30],
                       validations={"Status": STATUS_OPTS})
    for rr in range(r + 1, last + 1):
        ws[f"J{rr}"].fill = INPUT_FILL
        ws[f"K{rr}"].number_format = "mmm d, yyyy"
    ws.freeze_panes = ws.cell(row=r + 1, column=5)
    # the Overview COUNTIF uses whole column B: the header row 'Sprint' and title rows don't match S0..S6

    ws = wb.create_sheet("Reusable Blocks")
    brand_sheet(ws, title)
    r = title_block(ws, "Reusable blocks", "Built the same way in several flows. The steps say 'add the … block'.")
    rows = []
    for key, (name, items) in D.BLOCKS.items():
        for label, text in items:
            rows.append([name, label, text])
    write_table(ws, r, ["Block", "Action / field", "Value / expression"], rows, "tblBlocks", widths=[34, 70, 110])

    ws = wb.create_sheet("Gates")
    brand_sheet(ws, title)
    r = title_block(ws, "Gates", "Every line must be Passed before the next sprint starts.")
    rows = [[s[0], s[6], g, "Not started", None, ""] for s in D.SPRINTS for g in s[7]]
    last = write_table(ws, r, ["Sprint", "Gate", "Criterion", "Status", "Date", "Evidence / notes"], rows, "tblGates",
                       widths=[8, 26, 80, 13, 12, 50], validations={"Status": ["Not started", "Passed", "Failed", "Waived"]})
    for rr in range(r + 1, last + 1):
        ws[f"D{rr}"].fill = INPUT_FILL
        ws[f"E{rr}"].number_format = "mmm d, yyyy"

    ws = wb.create_sheet("Test Pack")
    brand_sheet(ws, title)
    r = title_block(ws, "Test Pack (UAT)", "Run in the sprint shown. Record the result.")
    rows = [list(t) + ["Not run", "", None, ""] for t in D.TESTS]
    last = write_table(ws, r, ["ID", "Sprint", "Scenario", "How", "Expected", "Result", "Tested by", "Date", "Notes"], rows, "tblTests",
                       widths=[7, 7, 30, 60, 80, 10, 14, 12, 30], validations={"Result": ["Not run", "Pass", "Fail"]})
    for rr in range(r + 1, last + 1):
        ws[f"F{rr}"].fill = INPUT_FILL
        ws[f"H{rr}"].number_format = "mmm d, yyyy"
    ws.freeze_panes = ws.cell(row=r + 1, column=4)

    for name, hdrs, data, tname, widths in [
        ("Decisions", ["ID", "Topic", "Decision", "Why", "Status"], D.DECISIONS, "tblDecisions", [6, 24, 80, 70, 16]),
        ("Inspiration Review", ["Idea", "Verdict", "Why", "How it shows up"], D.INSPIRATION, "tblInspiration", [50, 16, 60, 70]),
        ("Plan Traceability", ["Project plan item", "Where it lives in the build", "Sprint"], D.TRACE, "tblTrace", [62, 80, 10]),
        ("Dependencies", ["Item", "Owner", "Status", "Needed by", "Notes"], D.DEPENDENCIES, "tblDependencies", [60, 18, 14, 10, 60]),
    ]:
        ws = wb.create_sheet(name)
        brand_sheet(ws, title)
        r = title_block(ws, name, {"Decisions": "Design decisions and the ones that need your confirmation.",
                                   "Inspiration Review": "What we adopted, adapted, deferred or rejected from the email-ticketing inspiration — and why.",
                                   "Plan Traceability": "Every element of the project plan and where it is delivered.",
                                   "Dependencies": "From the plan's dependencies table, plus what this build adds."}[name])
        write_table(ws, r, hdrs, [list(x) for x in data], tname, widths=widths,
                    validations={"Status": ["To confirm", "Confirmed", "Ready (this kit)", "To build", "Pending", "Done"]} if name == "Dependencies" else None)

    ws = wb.create_sheet("Comms")
    brand_sheet(ws, title)
    r = title_block(ws, "Launch comms (Sprint 2)", "Replace <IntakeFormUrl> with the real link.")
    write_table(ws, r, ["Where", "Text"], [["#lane-people-talent announcement (pin it)", D.ANNOUNCEMENT], ["HR mailbox automatic reply", D.AUTO_REPLY]],
                "tblComms", widths=[40, 140])

    path = os.path.join(OUT, "HR Change Build Blueprint.xlsx")
    wb.save(path)
    return path


# ====================================================================== 3. CSV for the one list
def build_csv():
    cols = [f[0] for f in D.FIELDS if f[2] != "Person"]
    path = os.path.join(OUT, "HR-Change-Requests.csv")
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerow([D.CSV_SAMPLE[c] for c in cols])
    return path


# ====================================================================== 4. Word blueprint
def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_font(run, name="Rubik", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def fix_widths(t, widths):
    """Fixed layout + explicit grid so Word and LibreOffice honour the column widths."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    st = doc.styles["Normal"]
    st.font.name = "Rubik"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Rubik")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15
    for lvl, (font, size, color, bold, italic) in {1: ("Asap Condensed", 22, PLUM, True, True), 2: ("Comfortaa", 14, FOREST, True, False),
                                                    3: ("Comfortaa", 11.5, PLUM, True, False)}.items():
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.bold, h.font.italic, h.font.color.rgb = font, Pt(size), bold, italic, rgb(color)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True

    hp = sec.header.paragraphs[0]
    set_font(hp.add_run("Greater Good Health  |  HR Change + Orchestration — Build Blueprint"), size=8.5, color=FOREST)
    fp = sec.footer.paragraphs[0]
    set_font(fp.add_run("People Ops · internal · v2.0"), size=8, color=CHAR)

    def para(text="", bold=False, italic=False, color=CHAR, size=None, style=None):
        p = doc.add_paragraph(style=style)
        if text:
            set_font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(text), color=CHAR)
        return p

    def table(headers, rows, widths=None, head_fill=FOREST, font_size=9, header=True):
        t = doc.add_table(rows=1 if header else 0, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers if header else []):
            c = t.rows[0].cells[i]
            c.text = ""
            set_font(c.paragraphs[0].add_run(h), size=font_size, bold=True, color="FFFFFF")
            shade(c, head_fill)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                set_font(cells[i].paragraphs[0].add_run(str(v)), size=font_size, color=CHAR)
                if ri % 2 == 1:
                    shade(cells[i], MINT_LIGHT)
        if widths:
            fix_widths(t, widths)
        doc.add_paragraph()
        return t

    def paste_table(paste):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for label, text in paste:
            cells = t.add_row().cells
            set_font(cells[0].paragraphs[0].add_run(label), size=8.5, bold=True, color=FOREST)
            set_font(cells[1].paragraphs[0].add_run(text), size=8.5, color=CHAR)
            shade(cells[1], MINT_LIGHT)
        fix_widths(t, [5.2, 11.8])
        doc.add_paragraph()

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- cover
    for _ in range(5):
        doc.add_paragraph()
    para("HR Change + Orchestration", bold=True, italic=True, color=PLUM, size=34).runs[0].font.name = "Asap Condensed"
    para("Build Blueprint", color=FOREST, size=20).runs[0].font.name = "Comfortaa"
    para("One SharePoint list · Excel reference tables · Approvals · Slack visibility", color=CHAR, size=12)
    doc.add_paragraph()
    table(["", ""], [["Owner", "CJ"], ["Primary executor", "Johanna"], ["Reference", "Project 10k: People Team Owned Automations — Steps 5, 5b, 6"],
                     ["Version", "2.0 — replaces the Slack List design and the five-list v1 build"], ["Status", "Ready to start Sprint 0"]], widths=[4.5, 12.5], header=False)
    page_break()

    # ---- 1 how to use
    doc.add_heading("1. How to use this blueprint", 1)
    para("This is the build, in order. Each sprint has a goal, steps and a gate. Do the steps, run the tests, pass the gate, then move on. "
         "Every step has an ID (S1-03) that matches a row in the tracker workbook, where you set its status.")
    table(["File", "What it is", "Where it goes"], [
        ["HR Change Blueprint.docx", "This document — the step-by-step build", "Wherever you read it"],
        ["HR Change Build Blueprint.xlsx", "Tracker: roadmap with dates, every step with a status, gates, test pack, decisions", "Your working files"],
        [D.REF_FILE, "The reference tables the flows read: Settings, Change Types, Checklists, Approvers (+ list field reference)", "People Ops site → Documents → HR Change"],
        ["HR-Change-Requests.csv", "Creates the one SharePoint list", "Upload once in Sprint 0"],
    ], widths=[5, 7.5, 4.5])
    doc.add_heading("Conventions in every flow step", 3)
    for b in [
        "Rename every action to the name shown in bold. Expressions refer to actions by name with spaces replaced by _ (an action named 'Request ID' is outputs('Request_ID')).",
        "(fx) means: click in the field → fx (Insert expression) → paste → Add. Don't paste a leading @.",
        "Text containing @{ … }: type the plain text, and insert each @{ … } part with fx at that spot.",
        "SharePoint 'Update item' asks for required columns: pick the same field from the trigger (it writes back what's already there).",
        "Test loop: Save → Test → Manually → create or edit the item → open the run. SharePoint triggers check about once a minute.",
        "Build every action inside the 'Try' scope; the 'Catch' scope below it alerts you if anything fails (Appendix C).",
    ]:
        bullet(b)

    # ---- 2 architecture
    page_break()
    doc.add_heading("2. What we're building", 1)
    para("The project plan, rebuilt on Microsoft 365: the HR Change Slack List becomes ONE SharePoint list; the master templates, SLAs and approval "
         "matrix become tables in one Excel workbook; Power Automate does the orchestration; Slack stays where people see what's happening.")
    table(["Layer", "Plan (Slack design)", "This build"], [
        ["Front door", "Slack Workflow Builder form", "List New form, pinned in #lane-people-talent (D3)"],
        ["System of record", "HR Change Slack List", f"SharePoint list '{D.LIST_NAME}' — the only list"],
        ["Templates / SOP", "Master template items on the list", "Checklists table in the reference workbook"],
        ["Checklist delivery", "Message in the item thread", "Item comment (thread) + editable Checklist field with progress %"],
        ["Approvals", "Text steps / Slack DMs", "Power Automate Approvals, sequential, from the approval chain table"],
        ["Notifications", "Workflow Builder messages", "Slack connector posts + one daily digest"],
    ], widths=[3.2, 5.8, 8])
    doc.add_heading("The seven flows", 2)
    table(["Flow", "Plan reference", "Trigger", "What it does", "Sprint"], [
        ["1 Intake Notification", "Flow 1", "Item created", "Request ID + title, SLA due date, pay period/cutoff flag (I), default assignee, 📥 comment, People Ops ping, confirmation email", "S1"],
        ["2 Checklist Router", "Flow 2 (10 paths + Other)", "Item created", "Reads this type's steps from Excel, writes the Checklist field, posts the checklist comment; Other → 'needs manual scoping'", "S1"],
        ["3 Payroll Notification", "Flow 3 (Template I only)", "Item modified → I resolved", "One-way FYI to the Payroll room; stamps Payroll Notified On", "S2"],
        ["4 Completion Auto-Comment", "Completion auto-comment", "Item modified → 3-Resolved", "Stamps Completion Date, ticks 'resolution' steps, ✅ comment, close-the-loop email, optional People Ops summary", "S1"],
        ["5 Approval Routing", "Approval matrix (B, C, F, I)", "Item created", "Sequential approvals from the chain; history + comments; unlocks 🔒 steps or closes as Rejected", "S2–S3"],
        ["6 Checklist Progress", "Phase 2 progress bar", "Item modified", "Checklist % and Next Action from the [ ]/[x] ticks; 'ready to resolve' ping", "S3"],
        ["7 Daily Digest + Cutoff Reminder", "Nice-to-add + Phase 3 dashboard", "Daily 8:00", "SLA / next-action / exceptions digest; extra-shift reminder before each payroll cutoff", "S4"],
    ], widths=[3.4, 3.2, 2.8, 6.4, 1.2], font_size=8.5)
    para("Why separate flows: as the plan says for Flow 2, a failure in one shouldn't take down another — and each one is small enough to build and test in a sitting.", italic=True)

    # ---- 3 decisions
    doc.add_heading("3. Design decisions", 1)
    para("Three need confirmation before Sprint 0 ends (Status = Confirm).")
    table(["ID", "Topic", "Decision", "Status"], [[d[0], d[1], d[2], d[4]] for d in D.DECISIONS], widths=[1.2, 3.3, 10.2, 2.3], font_size=8.5)

    # ---- 4 inspiration
    doc.add_heading("4. What we took from the inspiration — and what we didn't", 1)
    para("The email-ticketing design was reviewed against the project plan's principles. The test: does it make the HR change process clearer "
         "or more reliable without breaking 'one front door'?")
    table(["Idea", "Verdict", "How it shows up (or why not)"], [[i[0], i[1], i[3] if i[1].startswith("Adopt") or i[1] == "Adapt" else i[2]] for i in D.INSPIRATION],
          widths=[6, 2.2, 8.8], font_size=8.5)

    # ---- 5 roadmap
    page_break()
    doc.add_heading("5. Roadmap — sprints and gates", 1)
    para("Two-week sprints (Sprint 0 is one week). Dates live in the tracker: set the start date on its Overview tab.")
    table(["Sprint", "Name", "Phase", "Weeks", "Gate"], [[s[0], s[1], s[3], s[2] or "ongoing", s[6]] for s in D.SPRINTS], widths=[1.5, 7.5, 1.8, 1.6, 4.6])
    for s in D.SPRINTS:
        doc.add_heading(f"{s[0]} gate — {s[6]}", 3)
        for g in s[7]:
            bullet("☐ " + g)

    # ---- sprints
    by_sprint = {}
    for stp in D.STEPS:
        by_sprint.setdefault(stp["sprint"], []).append(stp)
    test_by_sprint = {}
    for t in D.TESTS:
        test_by_sprint.setdefault(t[1], []).append(t)
    for n, s in enumerate(D.SPRINTS, start=6):
        page_break()
        doc.add_heading(f"{n}. {s[0]} — {s[1]}", 1)
        table(["Goal", "Entry criteria", "Phase"], [[s[4], s[5], s[3]]], widths=[9.5, 5.5, 2])
        for stp in by_sprint.get(s[0], []):
            doc.add_heading(f"STEP {stp['id']}: {stp['title']}", 3)
            if stp["where"]:
                p = para()
                set_font(p.add_run("Where: "), bold=True, color=FOREST)
                set_font(p.add_run(stp["where"]), color=CHAR)
            for d_ in stp["do"]:
                bullet(d_)
            if stp["paste"]:
                paste_table(stp["paste"])
            if stp["check"]:
                p = para()
                set_font(p.add_run("✔ Check: "), bold=True, color=DAHLIA)
                set_font(p.add_run(stp["check"]), color=CHAR)
        if test_by_sprint.get(s[0]):
            doc.add_heading(f"{s[0]} tests", 2)
            table(["ID", "Scenario", "Expected"], [[t[0], t[2], t[4]] for t in test_by_sprint[s[0]]], widths=[1.3, 4.2, 11.5], font_size=8.5)
        doc.add_heading(f"Gate — {s[6]}", 2)
        for g in s[7]:
            bullet("☐ " + g)

    # ---- appendices
    page_break()
    doc.add_heading("Appendix A–C — Reusable blocks", 1)
    for key, (name, items) in D.BLOCKS.items():
        doc.add_heading(name, 2)
        paste_table(items)
    doc.add_heading("Appendix D — Launch comms", 1)
    para("#lane-people-talent announcement (pin it):", bold=True, color=FOREST)
    para(D.ANNOUNCEMENT)
    para("HR mailbox automatic reply:", bold=True, color=FOREST)
    para(D.AUTO_REPLY)
    page_break()
    doc.add_heading("Appendix E — Test pack", 1)
    table(["ID", "Sprint", "Scenario", "How", "Expected"], [list(t) for t in D.TESTS], widths=[1.1, 1.2, 3.2, 5, 6.5], font_size=8)
    doc.add_heading("Appendix F — Plan traceability", 1)
    table(["Project plan item", "Where it lives in the build", "Sprint"], [list(t) for t in D.TRACE], widths=[7, 8.2, 1.8], font_size=8.5)
    doc.add_heading("Appendix G — Dependencies", 1)
    table(["Item", "Owner", "Status", "Needed by"], [list(x[:4]) for x in D.DEPENDENCIES], widths=[9, 3, 3, 2])
    doc.add_heading(f"Appendix H — The one list: {D.LIST_NAME}", 1)
    para("Full detail (choices, show formulas) is on the List Fields tab of the reference workbook.", italic=True)
    table(["Internal name", "Type", "Req.", "Group", "Show for", "Set by"], [[f[0], f[2], f[3], f[5], f[6], f[7]] for f in D.FIELDS],
          widths=[3.4, 3.2, 1, 1.9, 1.6, 5.9], font_size=8)

    path = os.path.join(OUT, "HR Change Blueprint.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    for fn in (build_reference, build_tracker, build_csv, build_docx):
        print("wrote", fn())
